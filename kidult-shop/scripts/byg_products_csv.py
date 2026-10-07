#!/usr/bin/env python3
"""Bygger products.csv til Shopify-import (Produkter → Importér).

Kør fra kidult-shop/:
    python3 scripts/byg_products_csv.py

Med kostpriser fra regnearket (kræver openpyxl). Filen holdes ude af git,
fordi repoet er offentligt:
    python3 scripts/byg_products_csv.py --kostpris sti/til/kidult-marginer.xlsx

Produktdata (dele, mål, byggetid) er hentet fra producenter og forhandlere.
Kilderne står i docs/produkter.md. Felter markeret TJEK der skal bekræftes,
før varen publiceres.
"""
import argparse
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

# Rækkefølgen og titlerne i "navn_i_ark" matcher arket Produkter i kidult-marginer.xlsx.
PRODUKTER = [
    {
        "navn_i_ark": "Golden Gate Bridge – GOLD (MMS001G)",
        "handle": "golden-gate-bridge-guld-metalmodel",
        "titel": "Golden Gate Bridge i guld – 3D-metalmodel",
        "mærke": "Metal Earth",
        "type": "3D-metalmodel",
        "sku": "BKC-ME-MMS001G",
        "pris": 89,
        "tags": ["Tilkøb", "Sværhed: Let", "Tema: Bygninger og vartegn"],
        "intro": "Verdens måske mest fotograferede bro, i en guldfarvet udgave på størrelse med en kuglepen. "
        "Delene sidder i et tyndt stålark, og du knækker dem forsigtigt ud og samler dem ved at bukke små tapper "
        "gennem hullerne. Der skal hverken lim eller lodning til, men en spidstang eller pincet gør det meget nemmere. "
        "Den færdige bro er fin på skrivebordet eller i reolen.",
        "fakta": [
            ("Byggetid", "Oplyses ikke af producenten. Ét ark og sværhedsgrad let gør den til et kort projekt"),
            ("Antal dele", "1 metalark (ca. 10 × 10 cm)"),
            ("Mål (færdig)", "ca. 15 × 0,8 × 4,5 cm"),
            ("Sværhedsgrad", "Let"),
            ("Anbefalet alder", "14+"),
        ],
        "hvem": "Til dig, der vil prøve metalmodeller for første gang, eller som en lille ekstra gave til en, der har stået på broen.",
    },
    {
        "navn_i_ark": "Wheel Loader – COLOR (MMS183)",
        "handle": "gummiged-metalmodel-farve",
        "titel": "Gummiged i farver – 3D-metalmodel",
        "mærke": "Metal Earth",
        "type": "3D-metalmodel",
        "sku": "BKC-ME-MMS183",
        "pris": 119,
        "tags": ["Tilkøb", "Sværhed: Udfordrende", "Tema: Køretøjer"],
        "intro": "En gummiged i orange og sort stål med skovl, førerhus og store hjul. Alle dele er forfarvede, så du skal "
        "ikke male noget. Du løsner delene fra arket og bukker tapperne på plads, uden lim. Modellen er klassificeret "
        "som udfordrende af producenten, fordi mange små dele skal bukkes præcist. Det er netop det, der gør den "
        "tilfredsstillende at få færdig.",
        "fakta": [
            ("Byggetid", "Oplyses ikke af producenten"),
            ("Antal dele", "108"),
            ("Mål (færdig)", "ca. 9,2 × 3,3 × 3,9 cm"),
            ("Sværhedsgrad", "Udfordrende"),
        ],
        "hvem": "Til den tålmodige, der kan lide maskiner og fine detaljer, og som har en pincet eller spidstang ved hånden.",
    },
    {
        "navn_i_ark": "Merry-Go-Round mini 3D puzzle (TG404)",
        "handle": "karrusel-mini-3d-traemodel",
        "titel": "Karrusel – mini 3D-træmodel",
        "mærke": "Rolife",
        "type": "3D-træmodel",
        "sku": "BKC-RL-TG404",
        "pris": 99,
        "tags": ["Tilkøb", "Sværhed: Mellem", "Tema: Forlystelser"],
        "intro": "En lille, klassisk karrusel i lasercut træ med heste, tag og alle de små pynte-detaljer. Delene er "
        "nummererede og klikkes sammen uden lim, så du kan bygge den ved spisebordet en stille aften. Den færdige "
        "karrusel er en hyggelig lille figur til vindueskarmen eller reolen, og en nem gave at sende med posten.",
        "fakta": [
            ("Byggetid", "Oplyses ikke af producenten"),
            ("Antal dele", "174"),
            ("Mål (færdig)", "ca. 12 × 12 × 15 cm"),
            ("Sværhedsgrad", "Mellem (3 ud af 5 hos producenten)"),
            ("Anbefalet alder", "8+"),
        ],
        "hvem": "Til dig, der vil prøve et træbyggesæt uden at gå all in, eller som en lille opmærksomhed til en, der elsker tivoli.",
    },
    {
        "navn_i_ark": "Compact Racer",
        "handle": "compact-racer-mekanisk-traebil",
        "titel": "Compact Racer – mekanisk træbil",
        "mærke": "Ugears",
        "type": "Mekanisk træmodel",
        "sku": "BKC-UG-122086",
        "pris": 149,
        "tags": ["Tilkøb", "Sværhed: Let", "Tema: Køretøjer"],
        "intro": "En lille racerbil i træ fra Ugears, som er kendt for mekaniske modeller, der samles helt uden lim. "
        "Delene trykkes ud af træpladerne og samles efter en tegnet vejledning. Compact Racer er et af de hurtigste "
        "Ugears-sæt at bygge og derfor et godt sted at starte, hvis du vil se, om mekaniske træmodeller er noget for dig.",
        "fakta": [
            ("Byggetid", "ca. 25–30 minutter"),
            ("Antal dele", "90"),
            ("Mål (færdig)", "ca. 13,3 × 3,3 × 6,2 cm"),
            ("Sværhedsgrad", "Let"),
        ],
        "hvem": "Til begynderen, til frokostpausen, eller som tilkøb til en, der allerede har fået smag for Ugears.",
    },
    {
        "navn_i_ark": "Wheel-Organizer 2.0",
        "handle": "wheel-organizer-mekanisk-penneholder",
        "titel": "Wheel-Organizer 2.0 – mekanisk penneholder",
        "mærke": "Ugears",
        "type": "Mekanisk træmodel",
        "sku": "BKC-UG-WHEELORG2",
        "pris": 149,
        "tags": ["Tilkøb", "Sværhed: Let", "Tema: Skrivebord", "Gaveidé"],
        "intro": "En penneholder til skrivebordet med en lille planetgear-mekanisme indeni. Når du drejer på hjulet, "
        "kører holderne rundt. Den har plads til seks kuglepenne, blyanter eller pensler. Som alle Ugears-modeller "
        "samles den af lasercut træ uden lim eller værktøj. Et byggesæt, du faktisk bruger hver dag bagefter, og som får folk til at spørge, hvor du har den fra.",
        "fakta": [
            ("Byggetid", "ca. 1 time"),
            ("Antal dele", "51"),
            ("Mål (færdig)", "ca. 11 × 9 × 9,3 cm"),
            ("Sværhedsgrad", "Let"),
            ("Anbefalet alder", "14+"),
        ],
        "hvem": "Til skrivebordsnørden, kollegaen eller den studerende. Både en lille gave og et lille projekt.",
    },
    {
        "navn_i_ark": "Tiny Stores – Pet Shop",
        "handle": "pet-shop-miniaturebutik",
        "titel": "Pet Shop – lille miniaturebutik",
        "mærke": "Rolife",
        "type": "Miniaturehus",
        "sku": "BKC-RL-TINYPET",
        "pris": 149,
        "tags": ["Tilkøb", "Sværhed: Let", "Tema: Bygninger og vartegn"],
        "intro": "En bitte lille dyrehandel med rød dør, gul facade og et skilt formet som et kødben. Indenfor venter en "
        "hvalp og en kat, foderposer og plejeartikler i miniformat. Delene er lavet af forfarvet træ og karton, så du "
        "ikke skal male, og de samles efter en billedvejledning, typisk uden lim og værktøj.",
        "fakta": [
            ("Byggetid", "ca. 1–1,5 time"),
            ("Antal dele", "TJEK hos leverandør"),
            ("Mål (færdig)", "ca. 11,4 × 10,2 × 4,6 cm (h × b × d)"),
            ("Sværhedsgrad", "Let"),
        ],
        "hvem": "Til dyreelskeren og til dig, der vil prøve miniaturer uden at give dig i kast med et helt hus.",
    },
    {
        "navn_i_ark": "Book nook – Secret Bloomy Nook (TGE04)",
        "handle": "secret-bloomy-nook-book-nook",
        "titel": "Secret Bloomy Nook – book nook",
        "mærke": "Rolife",
        "type": "Book nook",
        "sku": "BKC-RL-TGE04",
        "pris": 349,
        "tags": ["Kerne", "Sværhed: Let", "Tema: Book nook", "Gaveidé"],
        "intro": "En book nook er et lille diorama, der står mellem bøgerne i reolen, som et hemmeligt stræde ind i en anden "
        "verden. Secret Bloomy Nook er en blomsterfyldt gyde med små facader og detaljer, du bygger lag for lag. "
        "Producenten kalder den begyndervenlig, men med 246 dele og omkring seks timers byggetid er den et rigtigt "
        "aftenprojekt, gerne fordelt over et par gange.",
        "fakta": [
            ("Byggetid", "ca. 6 timer"),
            ("Antal dele", "246"),
            ("Mål (færdig)", "ca. 14 × 14,9 × 18 cm"),
            ("Sværhedsgrad", "Let (begynder), men lang byggetid"),
            ("Anbefalet alder", "14+"),
        ],
        "hvem": "Til bogelskeren, til dig, der vil have et skærmfrit projekt, og som gave til en, der har alt i forvejen.",
    },
    {
        "navn_i_ark": "Marble Explorer (LG503)",
        "handle": "marble-explorer-kuglebane-trae",
        "titel": "Marble Explorer – kuglebane i træ",
        "mærke": "ROKR",
        "type": "Mekanisk træmodel",
        "sku": "BKC-RK-LG503",
        "pris": 349,
        "tags": ["Kerne", "Sværhed: Mellem", "Tema: Kuglebane", "Gaveidé"],
        "intro": "En mekanisk kuglebane i lasercut krydsfiner, hvor stålkugler løftes op og triller ned ad baner, spiraler og "
        "trapper. Du bygger den af 260 dele uden lim, og bagefter er det svært at lade være med at sætte den i gang "
        "igen og igen. Byggeriet tager omkring seks timer, og du kan følge mekanikken tage form undervejs.",
        "fakta": [
            ("Byggetid", "ca. 6 timer"),
            ("Antal dele", "260"),
            ("Mål (færdig)", "ca. 25,5 × 22,9 × 20,4 cm"),
            ("Sværhedsgrad", "Mellem (4 ud af 6 hos producenten)"),
            ("Anbefalet alder", "14+"),
        ],
        "hvem": "Til dig, der kan lide at forstå, hvordan ting virker. Også en oplagt gave til en ingeniør eller en, der savner sin barndoms kuglebane.",
    },
    {
        "navn_i_ark": "Pink Beauty Secret miniature house",
        "handle": "pink-beauty-secret-miniaturehus",
        "titel": "Pink Beauty Secret – miniaturehus",
        "mærke": "Rolife",
        "type": "Miniaturehus",
        "sku": "BKC-RL-PINKBEAUTY",
        "pris": 379,
        "tags": ["Kerne", "Sværhed: Let", "Tema: Bygninger og vartegn", "Gaveidé"],
        "intro": "En lille skønhedsbutik i rosa nuancer med hylder, spejle og bittesmå produkter. Den er en del af Rolifes "
        "Super Creator-serie, hvor du bygger rummet op med rene linjer og mange små detaljer. Byggeriet tager omkring "
        "en time. Bemærk: tag og støvtæt dør sælges separat hos producenten, og den indbyggede LED skal bruge USB-C-strøm. "
        "Kabel og adapter følger ikke med.",
        "fakta": [
            ("Byggetid", "ca. 1 time"),
            ("Antal dele", "115"),
            ("Mål (færdig)", "ca. 16,3 × 16,3 × 15,2 cm"),
            ("Sværhedsgrad", "Let"),
        ],
        "hvem": "Til dig, der elsker miniaturer, rosa og hyggelige detaljer, og som en gave, der kan bygges på en enkelt aften.",
    },
    {
        "navn_i_ark": "Trimaran Merihobus",
        "handle": "trimaran-merihobus-mekanisk-sejlbaad",
        "titel": "Trimaran Merihobus – mekanisk sejlbåd i træ",
        "mærke": "Ugears",
        "type": "Mekanisk træmodel",
        "sku": "BKC-UG-70059",
        "pris": 599,
        "tags": ["Premium", "Sværhed: Udfordrende", "Tema: Skibe", "Gaveidé"],
        "intro": "En stor trimaran med tre skrog, master og sejl, bygget af 237 dele i lasercut træ uden lim eller "
        "specialværktøj. Med en højde på over en halv meter er den et samtalestykke på hylden eller i vindueskarmen. "
        "Byggetiden er 12–15 timer, så det er et projekt til flere weekender. Vejledningen er trin for trin og i farver.",
        "fakta": [
            ("Byggetid", "ca. 12–15 timer"),
            ("Antal dele", "237"),
            ("Mål (færdig)", "ca. 41 × 23,5 × 56 cm"),
            ("Sværhedsgrad", "Udfordrende"),
        ],
        "hvem": "Til sejleren, modelbyggeren og dig, der vil have et større projekt at se frem til i weekenden.",
    },
    {
        "navn_i_ark": "Dream Coffee Factory (EAB02)",
        "handle": "dream-coffee-factory-mekanisk-kaffefabrik",
        "titel": "Dream Coffee Factory – mekanisk kaffefabrik med musik",
        "mærke": "ROKR",
        "type": "Mekanisk træmodel",
        "sku": "BKC-RK-EAB02",
        "pris": 599,
        "tags": ["Premium", "Sværhed: Udfordrende", "Tema: Bygninger og vartegn", "Miniaturehus", "Gaveidé"],
        "intro": "En lille kaffefabrik i træ med bevægelige dele, LED-lys og spilledåse. Når den kører, bevæger mekanikken sig, "
        "og musikken spiller. Den samles med traditionelle tap-og-hul-samlinger uden lim, af 593 dele. Med omkring ti "
        "timers byggetid er den et af de største projekter i butikken. Den er både en mekanisk model og et miniaturehus.",
        "fakta": [
            ("Byggetid", "ca. 10 timer"),
            ("Antal dele", "593"),
            ("Mål (færdig)", "ca. 23,7 × 16,6 × 25 cm"),
            ("Sværhedsgrad", "Udfordrende"),
        ],
        "hvem": "Til kaffeelskeren og den erfarne bygger, eller som den store gave, der giver flere hyggelige aftener.",
    },
]

KOLONNER = [
    "Handle", "Title", "Body (HTML)", "Vendor", "Product Category", "Type", "Tags", "Published",
    "Option1 Name", "Option1 Value", "Variant SKU", "Variant Grams", "Variant Inventory Tracker",
    "Variant Inventory Qty", "Variant Inventory Policy", "Variant Fulfillment Service", "Variant Price",
    "Variant Compare At Price", "Variant Requires Shipping", "Variant Taxable", "Variant Barcode",
    "Image Src", "Image Position", "Image Alt Text", "Gift Card", "SEO Title", "SEO Description",
    "Variant Weight Unit", "Status",
]


def body_html(p):
    fakta = "".join(f"<li><strong>{k}:</strong> {v}</li>" for k, v in p["fakta"])
    return f"<p>{p['intro']}</p><ul>{fakta}</ul><p><strong>Hvem er den til?</strong> {p['hvem']}</p>"


def seo(p):
    """SEO-titel (maks. 60 tegn) og meta-beskrivelse (maks. 155 tegn) ud fra produktdata."""
    titel = f"{p['titel']} | Big Kid Club"
    if len(titel) > 60:
        titel = p["titel"][:60].rstrip(" –")
    fakta = dict(p["fakta"])
    detaljer = []
    if fakta.get("Antal dele", "").isdigit():
        detaljer.append(f"{fakta['Antal dele']} dele")
    if fakta.get("Byggetid", "").startswith("ca."):
        detaljer.append(f"byggetid {fakta['Byggetid']}")
    beskrivelse = f"{p['titel']} fra {p['mærke']}."
    if detaljer:
        beskrivelse += " " + ", ".join(detaljer).capitalize() + "."
    beskrivelse += " Byggesæt til voksne. Fri fragt fra 499 kr."
    return titel, beskrivelse[:155]


def række(p):
    seo_titel, seo_beskrivelse = seo(p)
    return {
        "Handle": p["handle"],
        "Title": p["titel"],
        "Body (HTML)": body_html(p),
        "Vendor": p["mærke"],
        "Product Category": "",
        "Type": p["type"],
        "Tags": ", ".join(p["tags"]),
        # Kladde uden lager, indtil du selv har bekræftet data, billeder og lager.
        "Published": "FALSE",
        "Option1 Name": "Title",
        "Option1 Value": "Default Title",
        "Variant SKU": p["sku"],
        "Variant Grams": "",
        "Variant Inventory Tracker": "shopify",
        "Variant Inventory Qty": "0",
        "Variant Inventory Policy": "deny",
        "Variant Fulfillment Service": "manual",
        "Variant Price": f"{p['pris']:.2f}",
        "Variant Compare At Price": "",
        "Variant Requires Shipping": "TRUE",
        "Variant Taxable": "TRUE",
        "Variant Barcode": "",
        "Image Src": "",
        "Image Position": "",
        "Image Alt Text": "",
        "Gift Card": "FALSE",
        "SEO Title": seo_titel,
        "SEO Description": seo_beskrivelse,
        "Variant Weight Unit": "g",
        "Status": "draft",
    }


def kostpriser(xlsx):
    import openpyxl

    ws = openpyxl.load_workbook(xlsx, data_only=True)["Produkter"]
    header = [c.value for c in ws[4]]
    navn, indkøb, fragt = (header.index(h) for h in ("Produkt", "Indkøb (kr)", "Indg. fragt (kr)"))
    # Varens kostpris = indkøb + indgående fragt. Betalingsgebyret hører til ordren, ikke varen.
    return {r[navn]: r[indkøb] + r[fragt] for r in ws.iter_rows(min_row=5, values_only=True)
            if r[navn] and isinstance(r[indkøb], (int, float))}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--kostpris", metavar="XLSX", help="kidult-marginer.xlsx; skriver products-med-kostpris.csv (ikke i git)")
    args = ap.parse_args()

    rækker = [række(p) for p in PRODUKTER]
    kolonner = KOLONNER
    ud = ROOT / "products.csv"

    if args.kostpris:
        kost = kostpriser(args.kostpris)
        mangler = [p["navn_i_ark"] for p in PRODUKTER if p["navn_i_ark"] not in kost]
        if mangler:
            sys.exit(f"Fandt ikke i arket: {', '.join(mangler)}")
        for r, p in zip(rækker, PRODUKTER):
            r["Cost per item"] = f"{kost[p['navn_i_ark']]:.2f}"
        kolonner = KOLONNER + ["Cost per item"]
        ud = ROOT / "products-med-kostpris.csv"

    with open(ud, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=kolonner)
        w.writeheader()
        w.writerows(rækker)

    for p in PRODUKTER:
        ord_ = len(" ".join([p["intro"], p["hvem"], *(f"{k} {v}" for k, v in p["fakta"])]).split())
        print(f"{p['sku']:<18} {p['pris']:>4} kr  {ord_:>3} ord  {p['titel']}")
    print(f"\nSkrev {len(rækker)} produkter til {ud.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
